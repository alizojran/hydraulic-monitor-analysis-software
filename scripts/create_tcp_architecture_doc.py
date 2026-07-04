from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "HMAS_TCP云端数据架构方案.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(92, 101, 115)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "C9D3DF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_paragraph_font(paragraph, size=11, color=INK, bold=False) -> None:
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D3DAE3", size="6")
    set_table_width(table, [6.35])
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(11)
    r.font.name = "Calibri"
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    for run in p2.runs:
        run.font.size = Pt(10.5)
        run.font.color.rgb = INK
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], LIGHT_GRAY)
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_font(p, size=10, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                if i > 0 and len(value) < 22:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_paragraph_font(p, size=9.5, color=INK, bold=False)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_code_block(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D7DDE5", size="4")
    set_table_width(table, [6.35])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for line_idx, line in enumerate(text.splitlines()):
        if line_idx:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(40, 45, 52)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_architecture_diagram(doc: Document) -> None:
    headers = ["现场侧", "上传链路", "阿里云侧", "前端侧"]
    rows = [
        [
            "CODESYS 控制器\n10 kHz 采集\n本地 FFT/特征提取",
            "TCP Client 主动连接\n长度头 + JSON 帧\n断线重连/seq",
            "TCP Server 接收\n鉴权/解析/入库\nWebSocket 实时转发",
            "本地 HTML/Vue\nHTTP 查询历史\nWebSocket 显示实时",
        ]
    ]
    table = add_table(doc, headers, rows, [1.55, 1.55, 1.65, 1.6])
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "HMAS TCP 云端数据架构方案"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run("Page ")
    r_begin = fp.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r_begin._r.append(fld_begin)
    r_instr = fp.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    r_instr._r.append(instr)
    r_end = fp.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_end._r.append(fld_end)
    for run in fp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def build_doc() -> None:
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("HMAS 液压监测系统")
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(23)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("CODESYS TCP 上传、阿里云数据库与本地 HTML 展示技术方案")
    sr.font.name = "Calibri"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.size = Pt(13)
    sr.font.color.rgb = MUTED

    meta_rows = [
        ("目标", "控制器内完成采集与 FFT，只上传频谱和特征值，不上传 10 kHz 原始波形"),
        ("通信", "CODESYS 作为 TCP Client 主动连接阿里云 TCP Server"),
        ("数据库", "推荐 PostgreSQL + TimescaleDB，分层保存 summary、spectrum、alarm"),
        ("版本日期", str(date.today())),
    ]
    add_table(doc, ["项目", "内容"], meta_rows, [1.25, 5.1])

    add_callout(
        doc,
        "结论",
        "推荐采用 CODESYS 主动 TCP 上传架构。PLC 端只做采集、FFT 和数据上传；阿里云端负责接收、鉴权、解析、入库和前端转发；本地 HTML/Vue 只通过 HTTP API 和 WebSocket 访问云端服务，不直接连接数据库。",
    )

    add_heading(doc, "1. 总体架构", 1)
    add_body(
        doc,
        "系统不再采用 MQTT，也不需要边缘网关。CODESYS 控制器主动连接阿里云服务器，上传 FFT 后的数据包。云端 TCP Server 接收数据后写入数据库，并通过 HTTP API 和 WebSocket 向本地 HTML 界面提供历史查询与实时显示能力。",
    )
    add_architecture_diagram(doc)

    add_heading(doc, "2. CODESYS 端职责", 1)
    add_body(
        doc,
        "CODESYS 端保留高频采集和本地频谱处理能力，但上传链路要保持简单、低优先级、可重连，避免影响采集任务实时性。",
    )
    for item in [
        "采集任务：高优先级，采集 10 kHz 原始数据并写入环形缓冲。",
        "FFT 任务：中优先级，每 1 秒取最近窗口数据，完成加窗、FFT、幅值 dB、RMS、Peak、Crest Factor、Top Peaks 和轴承特征频率幅值提取。",
        "上传任务：低优先级，维护 TCP 连接，发送 hello、heartbeat、fft、alarm 数据包，断线后定时重连。",
        "缓冲策略：周期性 FFT 数据允许少量丢包；告警事件应至少带 seq 并支持服务端去重。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. TCP 协议设计", 1)
    add_body(
        doc,
        "推荐采用固定长度头加 JSON 数据体。长度头解决 TCP 粘包和拆包问题，JSON 便于第一版调试和跨语言解析。后续如果频谱数组过大，可以把 spectrumDb 改为二进制 bytea 或 int16 压缩数组。",
    )
    add_code_block(
        doc,
        "[4字节大端长度][JSON payload]\n"
        "[4字节大端长度][JSON payload]\n"
        "[4字节大端长度][JSON payload]",
    )
    add_table(
        doc,
        ["包类型", "方向", "用途", "建议频率"],
        [
            ["hello", "PLC -> 云端", "设备上线、deviceId/secret 鉴权", "连接建立时"],
            ["heartbeat", "PLC -> 云端", "设备在线状态、版本、运行状态", "10 秒一次"],
            ["fft", "PLC -> 云端", "FFT 摘要、Top Peaks、轴承特征频率、可选完整频谱", "1-5 秒一次"],
            ["alarm", "PLC -> 云端", "告警事件、告警时刻频谱快照", "触发时立即发送"],
            ["ack", "云端 -> PLC", "确认已收到指定 seq", "可选，第一版可只记录 seq"],
        ],
        [1.0, 1.1, 3.2, 1.05],
    )

    add_heading(doc, "4. FFT 上传数据格式", 1)
    add_body(
        doc,
        "第一版建议每秒上传 summary，每 2-5 秒上传一次完整 spectrumDb。普通完整频谱只短期保存，报警频谱长期保存。",
    )
    add_code_block(
        doc,
        '{\n'
        '  "v": 1,\n'
        '  "type": "fft",\n'
        '  "deviceId": "HMAS-001",\n'
        '  "seq": 1024,\n'
        '  "ts": 1719300000000,\n'
        '  "channelId": "V02",\n'
        '  "sampleRate": 10000,\n'
        '  "fftSize": 2048,\n'
        '  "binHz": 4.8828125,\n'
        '  "rms": 0.42,\n'
        '  "peak": 1.16,\n'
        '  "crestFactor": 2.76,\n'
        '  "peaks": [{"f": 25.4, "db": -12.6}],\n'
        '  "bearing": {"bpfiHz": 135.2, "bpfiDb": -24.1, "bpfoHz": 89.7, "bpfoDb": -31.5},\n'
        '  "spectrumDb": [-80.1, -78.2, -75.4]\n'
        '}',
    )

    add_heading(doc, "5. 阿里云端服务", 1)
    add_body(
        doc,
        "云端至少包含 TCP 接收服务、数据库写入服务、HTTP 查询 API 和 WebSocket 实时推送服务。TCP Server 不建议直接把浏览器连接进来，浏览器仍然走标准 HTTP/WebSocket。",
    )
    for item in [
        "TCP Server：监听公网端口，接收 CODESYS 主动连接，解析长度头和 JSON payload。",
        "鉴权：hello 包校验 deviceId 和 secret；数据库只保存 secret_hash，避免明文密钥落库。",
        "入库：summary 长期保存，完整 spectrum 按保留策略保存，alarm 单独保存。",
        "前端接口：/api/latest 查询最新数据，/api/history 查询历史趋势，/ws/live 推送实时数据。",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "6. 数据库建议", 1)
    add_body(
        doc,
        "推荐 PostgreSQL + TimescaleDB。PostgreSQL 适合设备、通道、告警、配置等结构化数据，TimescaleDB 适合按时间写入和查询 FFT summary 与 spectrum。第一版可以在 ECS 自建，后续再迁移到 RDS PostgreSQL。",
    )
    add_table(
        doc,
        ["表", "保存内容", "保留周期", "说明"],
        [
            ["devices", "设备档案、secret_hash、last_seen_at", "长期", "设备注册与在线状态"],
            ["fft_summary", "RMS、Peak、Crest Factor、主频、轴承特征频率幅值、Top Peaks", "1-3 年", "趋势和报警分析主表"],
            ["fft_spectrum", "完整频谱数组 spectrumDb 或压缩 bytea", "7-30 天", "用于频谱回放和瀑布图"],
            ["alarm_events", "告警事件、告警时刻特征值、关联频谱", "长期", "维修追溯与诊断记录"],
        ],
        [1.2, 2.75, 1.1, 1.3],
    )

    add_heading(doc, "7. 数据量估算", 1)
    add_body(
        doc,
        "以下估算基于：V02 振动和 S01 声学两个通道，FFT size 为 2048，半谱 1024 点；summary 每 1 秒上传，完整频谱每 5 秒上传；原始 10 kHz 波形不上传。",
    )
    add_table(
        doc,
        ["数据类型", "上传频率", "网络数据量估算", "数据库占用估算"],
        [
            ["FFT summary / 特征值", "2 通道 x 每秒 1 次", "100-300 MB/天", "150-500 MB/天"],
            ["完整频谱 JSON", "2 通道 x 每 5 秒 1 次", "300-500 MB/天", "700 MB-1.5 GB/天"],
            ["完整频谱 int16/bytea", "2 通道 x 每 5 秒 1 次", "70-120 MB/天", "120-250 MB/天"],
            ["告警事件", "触发时", "可忽略", "可忽略"],
        ],
        [1.45, 1.65, 1.55, 1.55],
    )
    add_callout(
        doc,
        "容量建议",
        "第一版用 JSONB 保存完整频谱便于开发；链路稳定后将 spectrumDb 改为 int16 量化后写入 bytea，可明显降低网络与数据库占用。长期只保存 summary 和报警频谱，普通完整频谱保存 7-30 天。",
    )

    add_heading(doc, "8. 实施步骤", 1)
    for item in [
        "先做云端 TCP Server：实现长度头解析、hello 鉴权、fft 包入库、日志记录。",
        "在 CODESYS 中实现 TCP Client 上传任务：先上传模拟 FFT JSON，验证连接、重连和服务端解析。",
        "接入真实 FFT 结果：V02/S01 两个通道先跑通 summary，再加完整 spectrumDb。",
        "建立 PostgreSQL/TimescaleDB 表结构：先 JSONB，后续优化为 bytea 压缩。",
        "改造前端数据源：实时数据走云端 WebSocket，历史数据走 HTTP API。",
        "增加运行监控：记录设备 last_seen、seq 缺失、上传失败次数、数据库写入延迟。",
    ]:
        add_number(doc, item)

    add_heading(doc, "9. 风险与控制", 1)
    add_table(
        doc,
        ["风险", "影响", "控制措施"],
        [
            ["PLC 上传任务影响采集实时性", "采集抖动或丢样", "上传任务低优先级，发送缓冲限长，网络异常时丢弃普通周期包"],
            ["公网 TCP 连接不稳定", "实时数据中断", "定时重连，heartbeat，seq 缺号监控"],
            ["JSON 频谱数据过大", "带宽和磁盘增长快", "普通频谱降频上传，后续 int16 + bytea 压缩"],
            ["数据库无限增长", "磁盘和查询变慢", "TimescaleDB hypertable、保留策略、按设备和通道索引"],
            ["设备密钥泄露", "非法设备上传", "secret_hash、定期轮换、限制设备 IP 或增加签名字段"],
        ],
        [1.75, 1.45, 3.15],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
